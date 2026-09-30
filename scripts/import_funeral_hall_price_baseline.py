#!/usr/bin/env python3
import csv, io, json, os, re, sys, zipfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

DATASET_ID="15021763"
PAGE=f"https://www.data.go.kr/data/{DATASET_ID}/fileData.do"
OUT=Path("data/funeral-hall-prices.json")
LOCAL_CANDIDATES=[
 Path("data/source/2.장례식장가격정보_20230601.csv"),
 Path("data/source/2.장례식장가격정보_20230601(1).csv"),
 Path("2.장례식장가격정보_20230601.csv"),
 Path("2.장례식장가격정보_20230601(1).csv")
]
S=requests.Session()
S.headers.update({"User-Agent":"Mozilla/5.0 (compatible; dahamsangjo-publicdata-sync/1.0)","Accept-Language":"ko-KR,ko;q=0.9"})

def clean(v):
    return re.sub(r"\s+"," ",str(v or "").replace("\xa0"," ")).strip()

def norm(v):
    return re.sub(r"[^0-9A-Za-z가-힣]","",clean(v)).lower()

def fetch(url):
    r=S.get(url,timeout=45,allow_redirects=True)
    r.raise_for_status()
    return r

def discover_downloads():
    r=fetch(PAGE)
    soup=BeautifulSoup(r.text,"html.parser")
    urls=[]
    for a in soup.find_all("a",href=True):
        href=a["href"]
        txt=clean(a.get_text(" ",strip=True))
        if any(k in href.lower() for k in ["download","filedata","atch","file"]) or any(k in txt for k in ["다운로드","CSV","XLS","ZIP"]):
            u=urljoin(PAGE,href)
            if u not in urls: urls.append(u)
    # URLs embedded in onclick/javascript
    for m in re.finditer(r"https?://[^\"'<>\s]+|/[^\"'<>\s]*(?:download|file)[^\"'<>\s]*",r.text,re.I):
        u=urljoin(PAGE,m.group(0).replace("&amp;","&"))
        if "data.go.kr" in u and u not in urls: urls.append(u)
    return urls

def materialize_source():
    for p in LOCAL_CANDIDATES:
        if p.exists():
            return p.read_bytes(),p.name
    urls=discover_downloads()
    errors=[]
    for u in urls[:80]:
        try:
            r=fetch(u)
            cd=r.headers.get("content-disposition","")
            ct=r.headers.get("content-type","").lower()
            name=""
            m=re.search(r"filename\*?=(?:UTF-8''|\")?([^\";]+)",cd,re.I)
            if m:name=m.group(1)
            if not name:name=u.split("?")[0].split("/")[-1]
            b=r.content
            if len(b)<500: continue
            if b[:2]==b"PK" or "zip" in ct or name.lower().endswith(".zip"):
                with zipfile.ZipFile(io.BytesIO(b)) as z:
                    for zi in z.infolist():
                        n=zi.filename
                        if "장례식장" in n and ("가격" in n or n.lower().endswith(".csv")):
                            return z.read(zi),n
            if any(x in ct for x in ["csv","excel","spreadsheet","octet-stream"]) or name.lower().endswith((".csv",".xlsx",".xls")):
                if "장례식장" in name or "가격" in name:
                    return b,name
        except Exception as e:
            errors.append(f"{u}: {e}")
    # data.go.kr often builds file-download calls in script rather than href attributes.
    r=fetch(PAGE)
    text=r.text
    # expose hidden attachment IDs and ajax endpoints for data.go.kr's JS-driven downloader
    soup=BeautifulSoup(text,"html.parser")
    for inp in soup.find_all("input"):
        iid=inp.get("id","") or inp.get("name","")
        if "atch" in iid.lower() or "file" in iid.lower():
            print("DATA_GO_INPUT",iid,inp.get("value",""),file=sys.stderr)
    pats=[
        r".{0,300}limitAtchFileId.{0,700}",
        r".{0,300}atchFileId.{0,700}",
        r".{0,300}fileDownload.{0,700}",
        r".{0,300}download.{0,700}",
        r".{0,300}장례식장가격정보.{0,700}",
        r".{0,300}20230601.{0,700}"
    ]
    debug=[]
    for pat in pats:
        for m in re.finditer(pat,text,re.I|re.S):
            frag=m.group(0)
            if frag not in debug: debug.append(frag[:1600])
            if len(debug)>=20: break
        if len(debug)>=20: break
    for needle in ["limitAtchFileId","atchFileId","limitFileDetailSn"]:
        pos=text.find(needle)
        if pos>=0:
            frag=text[max(0,pos-1800):min(len(text),pos+2600)]
            if frag not in debug: debug.append(frag)
    print("DATA_GO_DEBUG_START",file=sys.stderr)
    for x in debug: print(x,file=sys.stderr)
    print("DATA_GO_DEBUG_END",file=sys.stderr)
    raise RuntimeError("official price file download link not resolved; candidates="+str(len(urls))+" errors="+str(errors[:3]))

def decode_csv(b):
    for enc in ["utf-8-sig","cp949","euc-kr","utf-8"]:
        try:return b.decode(enc)
        except UnicodeDecodeError:pass
    return b.decode("utf-8","replace")

def rows_from_bytes(b,name):
    ln=name.lower()
    if ln.endswith(".xlsx") or b[:2]==b"PK" and "csv" not in ln:
        wb=load_workbook(io.BytesIO(b),read_only=True,data_only=True)
        # choose sheet that mentions funeral halls/prices, else first
        ws=next((wb[s] for s in wb.sheetnames if "장례식장" in s or "가격" in s),wb[wb.sheetnames[0]])
        vals=list(ws.iter_rows(values_only=True))
        if not vals:return []
        headers=[clean(x) for x in vals[0]]
        return [dict(zip(headers,[clean(v) for v in row])) for row in vals[1:] if any(v not in (None,"") for v in row)]
    text=decode_csv(b)
    sample=text[:10000]
    try:dialect=csv.Sniffer().sniff(sample,delimiters=",\t|")
    except: dialect=csv.excel
    reader=csv.DictReader(io.StringIO(text),dialect=dialect)
    return [{clean(k):clean(v) for k,v in row.items() if k is not None} for row in reader]

def pick(row,*names):
    nmap={norm(k):v for k,v in row.items()}
    for name in names:
        nn=norm(name)
        if nn in nmap and clean(nmap[nn]):return clean(nmap[nn])
    for k,v in row.items():
        nk=norm(k)
        if any(norm(name) in nk for name in names) and clean(v):return clean(v)
    return ""

def amount_value(v):
    s=clean(v).replace(",","")
    m=re.search(r"(?<!\d)(\d{3,})(?!\d)",s)
    return int(m.group(1)) if m else 0

def classify(label):
    s=norm(label)
    if "빈소" in s or "분향" in s:return "빈소"
    if "안치" in s:return "안치실"
    if "염습" in s or "입관" in s:return "염습/입관"
    if "영결" in s:return "영결식장"
    if "청소" in s or "관리" in s:return "청소/관리"
    if "주차" in s:return "주차"
    if "도우미" in s:return "접객도우미"
    if "음식" in s or "식사" in s:return "음식"
    if "제단" in s:return "제단"
    return "기타"

def convert(rows,source_name):
    groups={}
    for row in rows:
        # tolerate multiple historical header spellings
        typ=pick(row,"시설구분","시설종류","장사시설구분","구분")
        if typ and "장례식장" not in typ and typ not in ("장례식장시설","장례식장"):
            # Some files are already funeral-hall-only and have a generic category column.
            if "장례식장" not in " ".join(row.keys()): pass
        name=pick(row,"시설명","장례식장명","업체명","상호")
        if not name: continue
        addr=pick(row,"주소","도로명주소","소재지")
        tel=pick(row,"전화번호","연락처","전화")
        label=pick(row,"품목명","항목명","가격항목","상품명","서비스명","세부품목","품목")
        spec=pick(row,"규격","내용","상세내용","비고")
        amount=0
        amount_raw=""
        for k,v in row.items():
            nk=norm(k)
            if any(x in nk for x in ["가격","금액","사용료","판매가"]) and "평균" not in nk:
                a=amount_value(v)
                if a>0:
                    amount=a; amount_raw=clean(v); break
        if not amount:
            # narrow fallback: price-looking numeric cell only when we have a label
            if label:
                for v in row.values():
                    a=amount_value(v)
                    if a>=1000: amount=a; amount_raw=clean(v); break
        if not label or not amount: continue
        key=(norm(name),norm(addr),norm(tel))
        rec=groups.setdefault(key,{
            "facilityName":name,"address":addr,"telephone":tel,
            "verifiedAt":"2023-06-02",
            "sourceName":"한국장례문화진흥원 전국 장사시설 현황 시설_가격정보",
            "sourceDatasetId":DATASET_ID,
            "sourceUrl":PAGE,
            "historical":True,
            "prices":[]
        })
        unit=""
        txt=" ".join([spec,amount_raw])
        m=re.search(r"(24시간|1일|시간당|1회|1실|\d+인분|kg|평)",txt)
        if m:unit=m.group(1)
        rec["prices"].append({"category":classify(label),"label":clean(label+" "+spec)[:180],"amount":amount,"unit":unit})

    out=[]
    for rec in groups.values():
        seen=set(); ps=[]
        for x in rec["prices"]:
            k=(x["label"],x["amount"],x["unit"])
            if k in seen:continue
            seen.add(k);ps.append(x)
        rec["prices"]=ps
        if ps:out.append(rec)
    return out

def main():
    b,name=materialize_source()
    rows=rows_from_bytes(b,name)
    print(f"source={name} rows={len(rows)}")
    items=convert(rows,name)
    if len(items)<10:
        # Never destroy the prior DB on parse/schema failure.
        print(f"WARN parsed only {len(items)} priced funeral halls; keeping existing database",file=sys.stderr)
        return 0
    out={
      "schemaVersion":3,
      "updatedAt":datetime.now(timezone.utc).isoformat(),
      "source":{
        "name":"한국장례문화진흥원 전국 장사시설 현황 시설_가격정보",
        "datasetId":DATASET_ID,
        "url":PAGE,
        "license":"공공누리 제1유형(출처표시)",
        "dataDate":"2023-06-02",
        "notice":"현재가격이 아닌 2023-06 기준 과거 공시가격. 실제 이용 전 e하늘/해당 장례식장 재확인 필요."
      },
      "pricedFacilities":len(items),
      "items":sorted(items,key=lambda x:(x.get("facilityName",""),x.get("address","")))
    }
    OUT.parent.mkdir(parents=True,exist_ok=True)
    tmp=OUT.with_suffix(".tmp")
    tmp.write_text(json.dumps(out,ensure_ascii=False,separators=(",",":")),encoding="utf-8")
    os.replace(tmp,OUT)
    print("saved",len(items),"facilities")
    return 0

if __name__=="__main__":
    raise SystemExit(main())
